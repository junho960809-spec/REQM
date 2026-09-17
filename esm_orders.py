"""ESM 원본 보존, 검증 및 판매전표 입력 변환. 네트워크/GUI와 분리한다."""
from __future__ import annotations

import hashlib
import json
import re
import shutil
import uuid
import zipfile
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import xlrd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

from ecount_sales_core import SmartStoreOrder, clean_identifier, parse_excel_date

STATUSES = ("입금대기", "배송대기", "배송준비", "배송중", "배송완료", "미수령신고", "정산예정", "정산완료")
DATE_TYPES = ("주문일", "결제완료일", "발송예정일", "발송일", "배송완료일", "정산완료일", "재배송일")
HEADERS = ("아이디", "상품명", "수량", "주문옵션", "판매금액", "판매자쿠폰할인", "개당 금액")
ESM_REQUIRED_HEADERS = {"아이디", "주문번호", "상품명", "수량", "주문옵션", "판매금액", "판매자쿠폰할인", "배송상태"}


def decimal_value(value, label: str) -> Decimal:
    try:
        result = Decimal(str(value).replace(",", "").strip())
        if not result.is_finite():
            raise InvalidOperation
        return result
    except (InvalidOperation, ValueError):
        raise ValueError(f"{label}: 숫자가 비어 있거나 올바르지 않습니다.") from None


def channel_for(account: str) -> str:
    if re.match(r"^[\[(]?\s*A(?:[\]()\s_]|$)", account, re.I) or "옥션" in account:
        return "옥션"
    if re.match(r"^[\[(]?\s*G(?:[\]()\s_]|$)", account, re.I) or "지마켓" in account or "G마켓" in account:
        return "지마켓"
    raise ValueError("아이디에서 옥션(A) 또는 지마켓(G)을 구분할 수 없습니다.")


@dataclass(frozen=True)
class EsmOrder:
    account: str
    product: str
    quantity: Decimal
    options: str
    sales: Decimal
    coupon: Decimal
    order_no: str
    product_no: str
    status: str
    ordered_at: datetime | None
    source_file: str
    source_row: int

    @property
    def net(self):
        return self.sales - self.coupon

    @property
    def unit(self):
        return self.net / self.quantity

    @property
    def key(self):
        # 상품/옵션이 다른 정상 주문행은 합치지 않는다.
        return (self.account, self.order_no, self.product_no, self.product, self.options)

    def voucher_order(self, voucher_date: date) -> SmartStoreOrder:
        identity = "ESM:" + hashlib.sha256(json.dumps(self.key, ensure_ascii=False).encode()).hexdigest()[:24]
        return SmartStoreOrder(
            source_row=self.source_row, order_no=f"{channel_for(self.account)}:{self.order_no}",
            product_order_no=identity, paid_at=self.ordered_at or datetime.combine(voucher_date, datetime.min.time()),
            status=self.status, product_name=self.product, options=self.options,
            quantity=self.quantity, item_total=self.net, initial_item_total=self.sales,
            final_item_total=self.net, source_type="ESM", include_shipping=False,
            source_channel=channel_for(self.account),
        )


def read_esm(path: Path | str) -> list[EsmOrder]:
    path = Path(path)
    signature = path.read_bytes()[:8]
    if signature.startswith(b"\xd0\xcf\x11\xe0"):
        book = xlrd.open_workbook(str(path))
        try:
            sheet = book.sheet_by_index(0)
            rows = [sheet.row_values(i) for i in range(sheet.nrows)]
        finally:
            book.release_resources()
    elif signature.startswith(b"PK"):
        # 파일 확장자와 달라도 실제 형식으로 읽는다.
        with path.open("rb") as stream:
            book = load_workbook(stream, read_only=True, data_only=True)
            try:
                rows = list(book.worksheets[0].values)
            finally:
                book.close()
    else:
        raise ValueError(f"{path.name}: XLS/XLSX 원본이 아닙니다. 로그인 페이지 또는 다운로드 오류인지 확인해주세요.")
    required = ESM_REQUIRED_HEADERS
    indexes = None
    for i, row in enumerate(rows[:30]):
        names = [str(v or "").strip().rstrip("*").strip() for v in row]
        if required.issubset(names):
            indexes = {name: names.index(name) for name in names if name}
            start = i + 1
            break
    if indexes is None:
        raise ValueError(f"{path.name}: ESM 필수 열(아이디·주문번호·상품명·수량·주문옵션·판매금액·판매자쿠폰할인·배송상태)이 없습니다.")
    result = []
    for number, row in enumerate(rows[start:], start=start + 1):
        if not any(v not in (None, "") for v in row):
            continue
        def value(name):
            c = indexes.get(name)
            return row[c] if c is not None and c < len(row) else None
        try:
            account = clean_identifier(value("아이디"))
            channel_for(account)
            quantity = decimal_value(value("수량"), "수량")
            sales = decimal_value(value("판매금액"), "판매금액")
            coupon = decimal_value(value("판매자쿠폰할인"), "판매자쿠폰할인")
            if quantity <= 0 or quantity != quantity.to_integral_value():
                raise ValueError("수량은 1 이상의 정수여야 합니다.")
            if sales < 0 or coupon < 0 or coupon > sales:
                raise ValueError("판매금액/쿠폰할인 범위를 확인해주세요.")
            order_no = clean_identifier(value("주문번호"))
            product = str(value("상품명") or "").strip()
            status = str(value("배송상태") or "").strip()
            if not order_no or not product:
                raise ValueError("주문번호 또는 상품명이 없습니다.")
            if status not in STATUSES:
                raise ValueError("수집 대상 8개 배송상태에 포함되지 않는 행입니다.")
            result.append(EsmOrder(account, product, quantity, str(value("주문옵션") or "").strip(),
                                   sales, coupon, order_no, clean_identifier(value("상품번호")), status,
                                   parse_excel_date(value("주문일자(결제확인전)")), path.name, number))
        except ValueError as exc:
            raise ValueError(f"{path.name} {number}행: {exc}") from None
    return result


def raw_esm_rows(path: Path | str) -> tuple[list[list], int]:
    """Return every original ESM row and the zero-based header row index."""
    path = Path(path)
    signature = path.read_bytes()[:8]
    if signature.startswith(b"\xd0\xcf\x11\xe0"):
        book = xlrd.open_workbook(str(path))
        try:
            sheet = book.sheet_by_index(0)
            rows = [sheet.row_values(i) for i in range(sheet.nrows)]
        finally:
            book.release_resources()
    elif signature.startswith(b"PK"):
        with path.open("rb") as stream:
            book = load_workbook(stream, read_only=True, data_only=False)
            try:
                rows = [list(row) for row in book.worksheets[0].values]
            finally:
                book.close()
    else:
        raise ValueError(f"{path.name}: XLS/XLSX 원본이 아닙니다.")
    for index, row in enumerate(rows[:30]):
        names = {str(value or "").strip().rstrip("*").strip() for value in row}
        if ESM_REQUIRED_HEADERS.issubset(names):
            return rows, index
    raise ValueError(f"{path.name}: ESM 원본 제목행을 찾지 못했습니다.")


def merge_orders(groups: list[list[EsmOrder]]) -> tuple[list[EsmOrder], int]:
    merged = {}
    duplicates = 0
    for group in groups:
        local = set()
        for order in group:
            if order.key in local:
                raise ValueError(f"{order.source_file} {order.source_row}행: 같은 주문 식별정보가 원본 안에 반복됩니다. 원본을 확인해주세요.")
            local.add(order.key)
            previous = merged.get(order.key)
            if previous:
                if (previous.quantity, previous.sales, previous.coupon, previous.status) != (order.quantity, order.sales, order.coupon, order.status):
                    raise ValueError("중복 주문의 상태·수량·금액이 다릅니다. 수집 중 변경되었을 수 있으므로 같은 기간을 다시 수집해주세요.")
                duplicates += 1
            merged[order.key] = order
    return list(merged.values()), duplicates


def file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class EsmSession:
    def __init__(self, folder: Path, manifest: dict):
        self.folder = folder
        self.manifest = manifest

    @classmethod
    def create(cls, root: Path, date_type: str, start: date, end: date, mode="자동 수집"):
        if date_type not in DATE_TYPES or end < start:
            raise ValueError("조회 기준 또는 조회 기간을 확인해주세요.")
        stamp = datetime.now()
        folder = root / stamp.strftime("%Y-%m-%d") / (stamp.strftime("%H%M%S") + "_" + uuid.uuid4().hex[:8])
        (folder / "원본").mkdir(parents=True)
        session = cls(folder, {"version": 1, "created_at": stamp.isoformat(timespec="seconds"), "mode": mode,
                              "date_type": date_type, "start_date": start.isoformat(), "end_date": end.isoformat(),
                              "site": "A/G 전체", "state": "진행 중", "files": [], "statuses": {}})
        session.save()
        return session

    @classmethod
    def open(cls, folder: Path):
        data = json.loads((folder / "수집기록.json").read_text(encoding="utf-8"))
        if data.get("version") != 1:
            raise ValueError("지원하지 않는 수집 기록입니다.")
        return cls(folder, data)

    def save(self):
        path = self.folder / "수집기록.json"
        temp = path.with_suffix(".tmp")
        temp.write_text(json.dumps(self.manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        temp.replace(path)

    def archive(self, source: Path, status: str, expected: int | None = None, original_name: str | None = None):
        suffix = ".xls" if source.read_bytes()[:8].startswith(b"\xd0\xcf") else ".xlsx"
        target = self.folder / "원본" / f"{len(self.manifest['files']) + 1:02d}_{status}{suffix}"
        shutil.copyfile(source, target)
        entry = {"file": target.relative_to(self.folder).as_posix(), "original_name": original_name or source.name,
                 "downloaded_at": datetime.now().isoformat(timespec="seconds"), "status": status,
                 "expected_count": expected, "sha256": file_hash(target)}
        self.manifest["files"].append(entry)
        self.save()  # 검증에 실패하더라도 원본은 남긴다.
        orders = read_esm(target)
        entry["row_count"] = len(orders)
        self.save()
        if expected is not None and len(orders) != expected:
            raise ValueError(f"{status}: 화면 {expected}건과 원본 {len(orders)}행이 다릅니다. 원본은 보관했습니다.")
        if status in STATUSES and any(o.status != status for o in orders):
            raise ValueError(f"{status}: 다운로드에 다른 상태의 주문이 포함되어 있습니다. 조회 조건을 확인해주세요.")
        self.manifest["statuses"][status] = {"state": "완료", "count": len(orders)}
        self.save()

    def empty(self, status):
        self.manifest["statuses"][status] = {"state": "완료", "count": 0}
        self.save()

    def orders(self):
        groups = []
        for entry in self.manifest["files"]:
            path = (self.folder / entry["file"]).resolve()
            if not path.is_relative_to((self.folder / "원본").resolve()) or path.suffix.lower() not in (".xls", ".xlsx"):
                raise ValueError("수집 기록의 원본 경로가 올바르지 않습니다.")
            if file_hash(path) != entry["sha256"]:
                raise ValueError("보관된 원본 파일이 변경되었습니다. 원본을 다시 가져와주세요.")
            rows = read_esm(path)
            if entry.get("expected_count") is not None and len(rows) != entry["expected_count"]:
                raise ValueError("조회 건수와 원본 행수가 다릅니다.")
            groups.append(rows)
        return merge_orders(groups)

    def finish(self):
        if self.manifest["mode"] == "자동 수집" and any(self.manifest["statuses"].get(s, {}).get("state") != "완료" for s in STATUSES):
            raise ValueError("8개 상태의 수집이 모두 완료되지 않았습니다.")
        orders, duplicates = self.orders()
        self.manifest.update(state="완료", order_count=len(orders), duplicate_count=duplicates)
        self.save()
        export_summary(self, self.folder / "ESM_취합.xlsx")
        if self.manifest["files"]:
            export_esm_original_format(self, self.folder / "ESM_원본양식_통합.xlsx")
        return orders

    def export_zip(self, destination: Path):
        files = [self.folder / "수집기록.json"]
        for entry in self.manifest["files"]:
            source = (self.folder / entry["file"]).resolve()
            if not source.is_relative_to((self.folder / "원본").resolve()) or file_hash(source) != entry["sha256"]:
                raise ValueError("원본 경로 또는 파일 검증에 실패했습니다.")
            files.append(source)
        if destination.resolve() in {p.resolve() for p in files}:
            raise ValueError("원본 파일을 덮어쓸 수 없습니다.")
        with zipfile.ZipFile(destination, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for path in files:
                archive.write(path, path.relative_to(self.folder))


def put_text(cell, value):
    cell.value = str(value)
    cell.data_type = "s"  # 상품명 등의 '='를 실행 수식으로 해석하지 않는다.


def export_summary(session: EsmSession, destination: Path):
    if session.manifest["state"] != "완료":
        raise ValueError("수집 완료 후 취합 파일을 저장할 수 있습니다.")
    orders, duplicates = session.orders()
    if destination.resolve().is_relative_to((session.folder / "원본").resolve()):
        raise ValueError("원본 폴더에는 취합 결과를 덮어쓸 수 없습니다.")
    book = Workbook()
    sheet = book.active
    sheet.title = "ESM 취합"
    sheet.append(HEADERS)
    for row, order in enumerate(orders, 2):
        for col, value in enumerate((order.account, order.product, order.quantity, order.options, order.sales, order.coupon), 1):
            if isinstance(value, str):
                put_text(sheet.cell(row, col), value)
            else:
                sheet.cell(row, col, value)
        sheet.cell(row, 7, f"=(E{row}-F{row})/C{row}")
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    for col, width in zip("ABCDEFG", (24, 62, 9, 38, 18, 20, 20)):
        sheet.column_dimensions[col].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="BDD7EE" if cell.column == 7 else "FFCC00")
        cell.font = Font(name="맑은 고딕", bold=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="맑은 고딕", size=10)
            cell.alignment = Alignment(vertical="center")
            if cell.column in (3, 5, 6, 7):
                cell.number_format = "#,##0.########"
    audit = book.create_sheet("수집 기록")
    for name, value in (("수집 시각", session.manifest["created_at"]), ("방식", session.manifest["mode"]),
                        ("조회 기준", session.manifest["date_type"]), ("조회 시작", session.manifest["start_date"]),
                        ("조회 종료", session.manifest["end_date"]), ("중복 제외 행수", duplicates)):
        audit.append([name, value])
    audit.append(["주문번호", "상품번호", "배송상태", "원본 파일", "원본 행"])
    for order in orders:
        r = audit.max_row + 1
        for c, value in enumerate((order.order_no, order.product_no, order.status, order.source_file, str(order.source_row)), 1):
            put_text(audit.cell(r, c), value)
    for col in "ABCDE":
        audit.column_dimensions[col].width = 30
    book.save(destination)
    book.close()


def export_esm_original_format(session: EsmSession, destination: Path):
    """Merge downloaded files into one workbook while retaining the ESM column layout.

    XLSX sources retain the first file's worksheet layout and cell styles. Legacy XLS
    sources retain the original preamble, complete column order and values in XLSX form.
    """
    if session.manifest["state"] != "완료":
        raise ValueError("수집 완료 후 ESM 원본 양식 통합 파일을 저장할 수 있습니다.")
    if not session.manifest["files"]:
        raise ValueError("통합할 ESM 원본 파일이 없습니다.")
    orders, _ = session.orders()
    selected = {(order.source_file, order.source_row) for order in orders}
    sources = [(session.folder / entry["file"]).resolve() for entry in session.manifest["files"]]
    first_rows, first_header = raw_esm_rows(sources[0])
    first_is_xlsx = sources[0].read_bytes()[:2] == b"PK"
    if first_is_xlsx:
        book = load_workbook(sources[0])
        sheet = book.worksheets[0]
        data_start = first_header + 2  # one-based first data row
        style_row = data_start if sheet.max_row >= data_start else first_header + 1
        styles = []
        for cell in sheet[style_row]:
            styles.append((copy(cell._style), copy(cell.number_format), copy(cell.alignment), copy(cell.protection)))
        if sheet.max_row >= data_start:
            sheet.delete_rows(data_start, sheet.max_row - data_start + 1)
    else:
        book = Workbook()
        sheet = book.active
        sheet.title = "ESM 주문"
        for row in first_rows[:first_header + 1]:
            sheet.append(row)
        data_start = first_header + 2
        styles = []
        for cell in sheet[first_header + 1]:
            cell.fill = PatternFill("solid", fgColor="FFCC00")
            cell.font = Font(name="맑은 고딕", bold=True)

    output_row = data_start
    for source in sources:
        rows, header = raw_esm_rows(source)
        if [str(v or "").strip() for v in rows[header]] != [str(v or "").strip() for v in first_rows[first_header]]:
            book.close()
            raise ValueError(f"{source.name}: 다른 ESM 열 구조가 포함되어 통합할 수 없습니다.")
        for source_row, values in enumerate(rows[header + 1:], start=header + 2):
            if (source.name, source_row) not in selected:
                continue
            for column, value in enumerate(values, 1):
                cell = sheet.cell(output_row, column, value)
                if column <= len(styles):
                    cell._style = copy(styles[column - 1][0])
                    cell.number_format = copy(styles[column - 1][1])
                    cell.alignment = copy(styles[column - 1][2])
                    cell.protection = copy(styles[column - 1][3])
            output_row += 1
    destination = Path(destination)
    if destination.resolve().is_relative_to((session.folder / "원본").resolve()):
        book.close()
        raise ValueError("원본 폴더에는 통합 결과를 덮어쓸 수 없습니다.")
    book.save(destination)
    book.close()
