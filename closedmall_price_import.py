from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook


CHANNEL_ALIASES = {
    "교보문고(핫트랙스)": "교보문고",
    "주식회사 현대이지웰": "이지웰",
    "삼성물산(주)패션부문": "SSF",
    "(주)이제너두": "이제너두",
    "(주)한섬": "한섬 EQL",
    "삼성카드 주식회사 쇼핑몰": "삼성카드 쇼핑몰",
    "삼성카드 주식회사 복지몰": "삼성카드 복지몰",
    "(주)현대홈쇼핑": "현대홈쇼핑(신)",
    "더블유컨셉코리아": "W컨셉",
    "주식회사 무신사": "무신사",
    "마켓컬리(주식회사 컬리)": "마켓컬리",
}


@dataclass(frozen=True)
class ClosedMallPriceRow:
    source_row: int
    customer_name: str
    source_channel: str
    customer_code: str
    product_name: str
    quantity: Decimal
    total: Decimal | None
    unit_price: Decimal | None
    status: str


def canonical_channel(customer_name: str) -> str:
    return CHANNEL_ALIASES.get(customer_name.strip(), customer_name.strip())


def read_closedmall_price_summary(path: str | Path) -> list[ClosedMallPriceRow]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook.active
        source_rows = sheet.iter_rows(min_col=1, max_col=7, values_only=True)
        next(source_rows, None)
        headers = list(next(source_rows, ()))
        required = ["거래처별", "품목별", "거래처코드", "수량", "공급가액", "부가세", "합계"]
        if headers != required:
            raise ValueError("폐쇄몰 판매현황집계 양식이 아닙니다. 2행의 열 제목을 확인해주세요.")

        result: list[ClosedMallPriceRow] = []
        for row_no, values in enumerate(source_rows, start=3):
            customer, product, customer_code, quantity, _supply, _vat, total = values
            if not customer or not product or customer_code in (None, ""):
                continue
            quantity_value = Decimal(str(quantity or 0))
            total_value = Decimal(str(total)) if total not in (None, "") else None
            unit_price = None
            status = "검수 필요"
            if "택배" in str(product) or "배송" in str(product):
                status = "배송비 제외"
            elif quantity_value <= 0 or total_value is None or total_value <= 0:
                status = "0원/미입력"
            else:
                calculated = total_value / quantity_value
                if calculated == calculated.to_integral_value():
                    unit_price = calculated
                    status = "매칭 필요"
                else:
                    status = "단가 검수"
            result.append(
                ClosedMallPriceRow(
                    source_row=row_no,
                    customer_name=str(customer).strip(),
                    source_channel=canonical_channel(str(customer)),
                    customer_code=str(customer_code).strip(),
                    product_name=str(product).strip(),
                    quantity=quantity_value,
                    total=total_value,
                    unit_price=unit_price,
                    status=status,
                )
            )
        return result
    finally:
        workbook.close()
