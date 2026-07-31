from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from excel_loader import B2C_PURCHASE_FORMAT, load_orders, missing_shipping_columns
from output_format_store import B2C_PURCHASE_FORMAT as B2C_PURCHASE_OUTPUT
from shipping_export import HEADERS as EXPORT_HEADERS, export_wekep, export_with_format


HEADERS = [
    "주문번호", "상품번호", "상품명", "옵션명", "수량", "판매단가", "판매금액",
    "수령자", "전화", "핸드폰", "우편번호", "주소", "배송메세지", "배송비", "송장출력갯수",
]


class B2CPurchaseLoaderTests(unittest.TestCase):
    def write_workbook(self, row: list[object]) -> Path:
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        path = Path(folder.name) / "b2c_purchase.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(HEADERS)
        sheet.append(row)
        workbook.save(path)
        workbook.close()
        return path

    def test_detects_purchase_format_without_order_number(self) -> None:
        path = self.write_workbook(
            ["", "", "REQM 상품", "블랙", 2, "", "", "홍길동", "", "01012345678",
             "01234", "서울시 중구 테스트로 1", "문 앞", "", ""]
        )

        orders, columns = load_orders(str(path))

        self.assertEqual(len(orders), 1)
        self.assertEqual(orders[0]["source_format"], B2C_PURCHASE_FORMAT)
        self.assertEqual(orders[0]["channel"], B2C_PURCHASE_FORMAT)
        self.assertEqual(orders[0]["order_number"], "")
        self.assertEqual(orders[0]["product_name"], "REQM 상품")
        self.assertEqual(orders[0]["quantity"], "2")
        self.assertEqual(orders[0]["recipient"], "홍길동")
        self.assertEqual(orders[0]["phone"], "01012345678")
        self.assertEqual(orders[0]["zipcode"], "01234")
        self.assertEqual(orders[0]["address"], "서울시 중구 테스트로 1")
        self.assertEqual(len(columns), 9)

        output = path.with_name("converted.xlsx")
        export_wekep(orders, str(output))
        workbook = load_workbook(output, data_only=True)
        try:
            sheet = workbook["택배출고"]
            self.assertEqual([cell.value for cell in sheet[1]], EXPORT_HEADERS)
            self.assertEqual(sheet["C2"].value, "REQM 상품")
            self.assertEqual(sheet["D2"].value, "2")
            self.assertEqual(sheet["E2"].value, "홍길동")
            self.assertEqual(sheet["F2"].value, "01012345678")
            self.assertEqual(sheet["G2"].value, "01234")
            self.assertEqual(sheet["H2"].value, "서울시 중구 테스트로 1")
        finally:
            workbook.close()

        purchase_output = path.with_name("purchase_converted.xlsx")
        export_with_format(orders, str(purchase_output), B2C_PURCHASE_OUTPUT)
        workbook = load_workbook(purchase_output, data_only=True)
        try:
            sheet = workbook["B2C 사입형"]
            self.assertEqual([cell.value for cell in sheet[1]], HEADERS)
            self.assertEqual(sheet["C2"].value, "REQM 상품")
            self.assertEqual(sheet["E2"].value, "2")
            self.assertEqual(sheet["H2"].value, "홍길동")
            self.assertEqual(sheet["J2"].value, "01012345678")
            self.assertEqual(sheet["K2"].value, "01234")
            self.assertEqual(sheet["L2"].value, "서울시 중구 테스트로 1")
        finally:
            workbook.close()

    def test_rejects_missing_colored_required_value(self) -> None:
        path = self.write_workbook(
            ["", "", "REQM 상품", "", 1, "", "", "홍길동", "", "",
             "01234", "서울시 중구 테스트로 1", "", "", ""]
        )

        with self.assertRaisesRegex(ValueError, r"2행 필수값 누락: 핸드폰\(J\)"):
            load_orders(str(path))

    def test_recipient_phone_number_alias_is_loaded(self) -> None:
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        path = Path(folder.name) / "seller_orders.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "주문번호", "상품명", "수량", "수령인", "수령인 전화번호",
            "우편번호", "주소",
        ])
        sheet.append([
            "ORDER-1", "REQM 상품", 1, "홍길동", "010-4966-0448",
            "01234", "서울시 중구 테스트로 1",
        ])
        workbook.save(path)
        workbook.close()

        orders, columns = load_orders(str(path))

        self.assertEqual(orders[0]["phone"], "010-4966-0448")
        self.assertEqual(missing_shipping_columns(columns), set())

    def test_missing_shipping_columns_are_reported_for_mapping_popup(self) -> None:
        columns = {
            "order_number": 0,
            "product_name": 1,
            "quantity": 2,
            "recipient": 3,
        }

        self.assertEqual(
            missing_shipping_columns(columns),
            {"phone", "zipcode", "address1"},
        )

    def test_custom_output_template_clears_examples_and_uses_selected_columns(self) -> None:
        path = self.write_workbook(
            ["", "", "REQM 상품", "", 1, "", "", "홍길동", "", "01012345678",
             "01234", "서울시 중구 테스트로 1", "", "", ""]
        )
        orders, _ = load_orders(str(path))
        template = path.with_name("custom_template.xlsx")
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["받는분", "제품", "개수", "주소지", "메모"])
        sheet.append(["예시 수령인", "예시 제품", 99, "예시 주소", "유지할 메모"])
        workbook.save(template)
        workbook.close()
        output = path.with_name("custom_output.xlsx")
        profile = {
            "id": "custom",
            "template_path": str(template),
            "header_row": 0,
            "mapping": {
                "recipient": "받는분",
                "product_name": "제품",
                "quantity": "개수",
                "address": "주소지",
            },
        }

        export_with_format(orders, str(output), profile)

        workbook = load_workbook(output, data_only=True)
        try:
            sheet = workbook.active
            self.assertEqual(sheet["A2"].value, "홍길동")
            self.assertEqual(sheet["B2"].value, "REQM 상품")
            self.assertEqual(sheet["C2"].value, "1")
            self.assertEqual(sheet["D2"].value, "서울시 중구 테스트로 1")
            self.assertEqual(sheet["E2"].value, "유지할 메모")
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
