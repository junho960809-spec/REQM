import json
import os
import zipfile
from dataclasses import replace
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from esm_orders import EsmSession, STATUSES, channel_for, read_esm, merge_orders, export_summary, export_esm_original_format
from esm_browser import browser_channels, date_windows, installed_browser, is_logged_in_url
from ecount_sales_core import ReferenceCatalog, convert_orders


def fixture_file(path, rows=None, quantity=5, coupon=15000):
    book = Workbook()
    sheet = book.active
    sheet.append(["메모"])
    # 열 순서가 달라도 제목으로 찾아야 한다.
    sheet.append(["주문번호*", "아이디*", "상품명", "수량", "판매자쿠폰할인", "판매금액", "주문옵션", "배송상태", "상품번호*", "주문일자(결제확인전)"])
    for row in rows or [["00001", "G(test)", "상품", quantity, coupon, "214,500", "블랙", "배송중", "00002", "2026-09-09"]]:
        sheet.append(row)
    book.save(path)
    book.close()
    return path


def test_read_real_format_numeric_strings_and_reordered_headers(tmp_path):
    rows = read_esm(fixture_file(tmp_path / "orders.xlsx"))
    assert rows[0].unit == Decimal("39900")
    assert rows[0].order_no == "00001"
    assert rows[0].product_no == "00002"
    assert rows[0].ordered_at == datetime(2026, 9, 9)
    assert channel_for("A(test)") == "옥션"
    assert channel_for("G_test") == "지마켓"


def test_esm_browser_prefers_chrome_before_edge():
    assert browser_channels() == ("chrome", "msedge")


def test_installed_browser_selects_chrome_when_available(monkeypatch, tmp_path):
    chrome = tmp_path / "chrome.exe"
    chrome.write_bytes(b"")
    monkeypatch.setenv("PROGRAMFILES", str(tmp_path))
    monkeypatch.setenv("PROGRAMFILES(X86)", str(tmp_path / "missing"))
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "missing"))
    monkeypatch.setattr("esm_browser.shutil.which", lambda name: str(chrome) if name == "chrome" else None)
    assert installed_browser() == ("chrome", str(chrome))


def test_esm_login_url_detection():
    assert is_logged_in_url("https://www.esmplus.com/Home/Home")
    assert is_logged_in_url("https://www.esmplus.com/Escrow/SmartDelivery/test")
    assert not is_logged_in_url("https://signin.esmplus.com/login")


@pytest.mark.parametrize("quantity,coupon", [(0, 0), ("", 0), ("nan", 0), (1.5, 0), (1, ""), (1, "bad"), (1, 999999), (1, -1)])
def test_invalid_values_do_not_become_zero(tmp_path, quantity, coupon):
    with pytest.raises(ValueError):
        read_esm(fixture_file(tmp_path / "bad.xlsx", quantity=quantity, coupon=coupon))


def test_dedup_keeps_different_items_and_blocks_conflicts(tmp_path):
    order = read_esm(fixture_file(tmp_path / "orders.xlsx"))[0]
    rows, duplicates = merge_orders([[order], [order]])
    assert len(rows) == 1 and duplicates == 1
    assert len(merge_orders([[order, replace(order, options="화이트")]])[0]) == 2
    assert len(merge_orders([[order, replace(order, account="A(test)")]])[0]) == 2
    with pytest.raises(ValueError, match="원본 안에 반복"):
        merge_orders([[order, order]])
    with pytest.raises(ValueError, match="다릅니다"):
        merge_orders([[order], [replace(order, status="정산예정")]])
    with pytest.raises(ValueError, match="다릅니다"):
        merge_orders([[order], [replace(order, coupon=Decimal(0))]])


def session_with_file(tmp_path):
    source = fixture_file(tmp_path / "source.xlsx")
    session = EsmSession.create(tmp_path / "archive", "주문일", date(2026, 9, 9), date(2026, 9, 9))
    session.archive(source, "배송중", 1)
    return source, session


def test_archive_exact_bytes_complete_gate_export_and_reopen(tmp_path):
    source, session = session_with_file(tmp_path)
    assert source.read_bytes() == (session.folder / session.manifest["files"][0]["file"]).read_bytes()
    with pytest.raises(ValueError, match="모두 완료"):
        session.finish()
    for status in STATUSES:
        if status != "배송중":
            session.empty(status)
    assert len(session.finish()) == 1
    restored = EsmSession.open(session.folder)
    assert len(restored.orders()[0]) == 1
    summary = load_workbook(session.folder / "ESM_취합.xlsx")
    assert summary.worksheets[0]["G2"].value == "=(E2-F2)/C2"
    assert summary.worksheets[0]["E2"].value == 214500
    assert summary.worksheets[0]["A1"].fill.fgColor.rgb.endswith("FFCC00")
    summary.close()
    raw_merged = load_workbook(session.folder / "ESM_원본양식_통합.xlsx")
    assert raw_merged.active["A3"].value == "00001"
    assert raw_merged.active["B3"].value == "G(test)"
    assert raw_merged.active.max_column == 10
    raw_merged.close()
    output = tmp_path / "raw.zip"
    session.export_zip(output)
    with zipfile.ZipFile(output) as z:
        assert z.read(session.manifest["files"][0]["file"]) == source.read_bytes()
        assert "수집기록.json" in z.namelist()
    original = session.folder / session.manifest["files"][0]["file"]
    original.write_bytes(b"tampered")
    with pytest.raises(ValueError, match="변경"):
        restored.orders()


def test_count_mismatch_preserves_original_and_rejects(tmp_path):
    source = fixture_file(tmp_path / "source.xlsx")
    session = EsmSession.create(tmp_path / "archive", "주문일", date(2026, 9, 9), date(2026, 9, 9))
    with pytest.raises(ValueError, match="다릅니다"):
        session.archive(source, "배송중", 2)
    assert len(session.manifest["files"]) == 1
    assert session.manifest["state"] != "완료"
    with pytest.raises(ValueError):
        export_summary(session, tmp_path / "no.xlsx")


def test_different_download_status_rejected(tmp_path):
    source = fixture_file(tmp_path / "source.xlsx")
    session = EsmSession.create(tmp_path / "archive", "주문일", date(2026, 9, 9), date(2026, 9, 9))
    with pytest.raises(ValueError, match="다른 상태"):
        session.archive(source, "배송준비", 1)


def test_original_format_merge_keeps_full_header_and_deduplicates(tmp_path):
    source = fixture_file(tmp_path / "source.xlsx")
    session = EsmSession.create(tmp_path / "archive", "주문일", date(2026, 9, 9), date(2026, 9, 9), "파일 가져오기")
    session.archive(source, "가져온원본1")
    session.archive(source, "가져온원본2")
    session.finish()
    destination = tmp_path / "merged.xlsx"
    export_esm_original_format(session, destination)
    book = load_workbook(destination)
    sheet = book.active
    assert sheet.max_row == 3
    assert sheet[2][0].value == "주문번호*"
    assert sheet[3][0].value == "00001"
    assert sheet.max_column == 10
    book.close()


def test_snapshot_folders_unique_and_period_split():
    windows = list(date_windows(date(2026, 8, 1), date(2026, 9, 9)))
    assert windows == [(date(2026, 8, 1), date(2026, 8, 31)), (date(2026, 9, 1), date(2026, 9, 9))]


def test_esm_conversion_separates_channels_and_preserves_net(tmp_path):
    order = read_esm(fixture_file(tmp_path / "orders.xlsx"))[0]
    channels = [{"source_name": c, "ecount_customer_code": code} for c, code in [("지마켓", "G-CUST"), ("옥션", "A-CUST")]]
    mappings = [{"mapping_key": c, "source_channel": c, "normalized_source": "상품블랙", "mapping_type": "single", "review_status": "confirmed"} for c in ("지마켓", "옥션")]
    components = [{"mapping_key": c, "sequence": 1, "item_code": "TEST", "quantity": 1} for c in ("지마켓", "옥션")]
    catalog = ReferenceCatalog([{"item_code": "TEST", "representative_name": "테스트"}], channels, mappings, components, [], [])
    orders = [order.voucher_order(date.today()), replace(order, account="A(test)").voucher_order(date.today())]
    result = convert_orders(orders, catalog)
    assert result.is_reconciled and not result.issues
    assert result.output_total == 399000
    assert {r.customer_code for r in result.lines} == {"G-CUST", "A-CUST"}
    assert result.shipping_total == 0
    assert len({o.product_order_no for o in orders}) == 2


def test_html_login_download_rejected(tmp_path):
    source = tmp_path / "error.xls"
    source.write_text("<html>login required</html>")
    with pytest.raises(ValueError, match="로그인"):
        read_esm(source)


def test_product_text_not_formula_and_source_zip_path_guard(tmp_path):
    source, session = session_with_file(tmp_path)
    book = load_workbook(source)
    book.active["C3"] = "=1+1"
    book.active["C3"].data_type = "s"
    book.save(source)
    book.close()
    other = EsmSession.create(tmp_path / "archive", "주문일", date.today(), date.today(), "파일 가져오기")
    assert other.folder != session.folder
    other.archive(source, "원본")
    other.finish()
    out = load_workbook(other.folder / "ESM_취합.xlsx")
    assert out.worksheets[0]["B2"].data_type == "s"
    out.close()
    other.manifest["files"][0]["file"] = "../source.xlsx"
    with pytest.raises(ValueError, match="경로"):
        other.orders()


def test_dialog_preview_and_transfer_buttons(tmp_path):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication, QDialog
    from PySide6.QtCore import QDate
    from esm_dialog import EsmSourceDialog
    app = QApplication.instance() or QApplication([])
    source, session = session_with_file(tmp_path)
    for status in STATUSES:
        if status != "배송중":
            session.empty(status)
    session.finish()
    dialog = EsmSourceDialog(QDate.currentDate(), QDate.currentDate(), session=session)
    assert dialog.table.item(0, 6).text() == "39,900"
    assert dialog.import_button.isEnabled()
    assert dialog.excel_button.isEnabled()
    dialog.set_busy(True)
    assert not dialog.import_button.isEnabled()
    assert not dialog.history_button.isEnabled()
    dialog.set_busy(False)
    dialog.use_orders()
    assert dialog.result() == QDialog.Accepted
